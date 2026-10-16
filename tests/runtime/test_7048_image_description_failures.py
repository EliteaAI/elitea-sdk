"""An image the loaders could not describe must reach the indexing report, not only the log.

Re-breaks if perform_llm_prediction_for_image_bytes stops recording its failures, if the indexer
stops collecting them per document, or if the report stops listing them.
"""

import io
from types import SimpleNamespace

import fitz
import pytest
from docx import Document as DocxDocument
from langchain_core.documents import Document
from PIL import Image, ImageDraw

from elitea_sdk.runtime.langchain.document_loaders.utils import (
    collect_image_description_failures,
    perform_llm_prediction_for_image_bytes,
)
from elitea_sdk.runtime.utils.utils import IndexerKeywords
from elitea_sdk.tools.base_indexer_toolkit import (
    DEPENDENT_DOC_META_KEY,
    BaseIndexerToolkit,
    IndexingStats,
    IndexingStatus,
)
from elitea_sdk.tools.utils.content_parser import process_content_by_type

BODY_TEXT = "Plain body text"
IMAGE_CONFIG = {
    ".pdf": {"extract_images": True, "use_llm": True, "use_default_prompt": True, "prompt": ""},
    ".docx": {"use_llm": True, "use_default_prompt": True, "prompt": ""},
}


class _RefusingLlm:
    def invoke(self, messages):
        raise RuntimeError("team not allowed to access model")


class _DescribingLlm:
    def invoke(self, messages):
        return SimpleNamespace(content="SECRETWORD")


def _png_bytes():
    image = Image.new("RGB", (400, 120), "white")
    ImageDraw.Draw(image).text((20, 40), "SECRETWORD", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _pdf_with_image():
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), BODY_TEXT)
    page.insert_image(fitz.Rect(72, 120, 472, 240), stream=_png_bytes())
    return pdf.tobytes()


def _docx_with_image():
    document = DocxDocument()
    document.add_paragraph(BODY_TEXT)
    document.add_picture(io.BytesIO(_png_bytes()))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _parse(content, filename, llm):
    return list(process_content_by_type(content, filename, llm=llm, chunking_config=IMAGE_CONFIG))


class TestCollector:
    def test_a_failed_description_is_recorded_and_still_raised(self):
        with collect_image_description_failures() as failures:
            with pytest.raises(RuntimeError):
                perform_llm_prediction_for_image_bytes(_png_bytes(), _RefusingLlm(), "Describe")

        assert failures == ["team not allowed to access model"]

    def test_a_successful_description_records_nothing(self):
        with collect_image_description_failures() as failures:
            perform_llm_prediction_for_image_bytes(_png_bytes(), _DescribingLlm(), "Describe")

        assert failures == []

    def test_outside_a_collector_the_failure_only_raises(self):
        with pytest.raises(RuntimeError):
            perform_llm_prediction_for_image_bytes(_png_bytes(), _RefusingLlm(), "Describe")


@pytest.mark.parametrize("filename,content", [
    pytest.param("report.pdf", _pdf_with_image(), id="pdf"),
    pytest.param("report.docx", _docx_with_image(), id="docx"),
])
class TestLoadersReportTheLostImage:
    def test_the_text_is_kept_and_the_image_failure_is_recorded(self, filename, content):
        with collect_image_description_failures() as failures:
            chunks = _parse(content, filename, _RefusingLlm())

        assert BODY_TEXT in " ".join(chunk.page_content for chunk in chunks)
        assert len(failures) == 1

    def test_a_described_image_records_nothing(self, filename, content):
        with collect_image_description_failures() as failures:
            chunks = _parse(content, filename, _DescribingLlm())

        assert "SECRETWORD" in " ".join(chunk.page_content for chunk in chunks)
        assert failures == []


class _Indexer:
    def __init__(self):
        self.stats = IndexingStats()

    def get_indexing_stats(self):
        return self.stats

    def _init_indexing_stats(self):
        return self.stats


class TestIndexerTracksLostImages:
    def test_each_lost_image_of_an_indexed_document_is_listed(self):
        indexer = _Indexer()

        BaseIndexerToolkit._parse_tracking_lost_images(
            indexer, "bucket/report.pdf", lambda: _parse(_pdf_with_image(), "report.pdf", _RefusingLlm())
        )

        assert indexer.stats.images_not_described == {"bucket/report.pdf: lost image 1"}

    def test_same_named_attachments_of_different_parents_are_counted_apart(self):
        indexer = _Indexer()

        for parent in ("page-1", "page-2"):
            label = BaseIndexerToolkit._lost_image_label({"parent_id": parent}, "spec.pdf", dependent=True)
            BaseIndexerToolkit._parse_tracking_lost_images(
                indexer, label, lambda: _parse(_pdf_with_image(), "spec.pdf", _RefusingLlm())
            )

        assert indexer.stats.images_not_described == {
            "spec.pdf (in page-1): lost image 1", "spec.pdf (in page-2): lost image 1",
        }

    def test_a_base_document_is_labelled_by_its_own_name(self):
        assert BaseIndexerToolkit._lost_image_label({"parent_id": "page-1"}, "report.pdf", dependent=False) == "report.pdf"

    def test_a_document_that_produced_nothing_is_not_counted_twice(self):
        indexer = _Indexer()

        def failed_parse():
            with pytest.raises(RuntimeError):
                perform_llm_prediction_for_image_bytes(_png_bytes(), _RefusingLlm(), "Describe")
            return []

        chunks = BaseIndexerToolkit._parse_tracking_lost_images(indexer, "photo.png", failed_parse)

        assert chunks == []
        assert indexer.stats.images_not_described == set()


class TestReport:
    def test_lost_images_are_a_failed_dependent_group_with_their_own_noun(self):
        stats = IndexingStats(images_not_described={"report.pdf: image 1", "report.pdf: image 2"})

        report = stats.build_report(
            status=IndexingStatus.OK, indexed_count=1,
            item_labels=("file", "files"), dependent_labels=("attachment", "attachments"),
        )

        failed = next(category for category in report["categories"] if category["kind"] == "failed")
        group = next(group for group in failed["groups"] if group["reason"] == "image_description_error")
        assert group["items"] == ["report.pdf: image 1", "report.pdf: image 2"]
        assert group["item_labels"] == {"singular": "image", "plural": "images"}
        assert group["dependent"] is True and group["counted"] is False
        assert failed["count"] == 0
        assert report["totals"]["dependent_not_indexed"] == 2
        assert report["totals"]["indexed"] == 1

    def test_the_text_summary_names_the_lost_images(self):
        stats = IndexingStats(images_not_described={"report.pdf: image 1"})

        assert "report.pdf: image 1" in stats.get_summary()


class TestIndexingPipeline:
    def test_a_lost_image_in_an_attachment_is_reported_with_its_parent(self):
        toolkit = BaseIndexerToolkit.model_construct(llm=_RefusingLlm(), embeddings=None)
        toolkit._indexing_stats = IndexingStats()
        attachment = Document(page_content="", metadata={
            IndexerKeywords.CONTENT_FILE_NAME.value: "spec.pdf",
            IndexerKeywords.CONTENT_IN_BYTES.value: _pdf_with_image(),
            IndexerKeywords.PARENT.value: "page-1",
            DEPENDENT_DOC_META_KEY: True,
        })

        chunks = list(toolkit._apply_loaders_chunkers(iter([attachment]), chunking_config=dict(IMAGE_CONFIG)))

        assert BODY_TEXT in " ".join(chunk.page_content for chunk in chunks)
        assert toolkit.get_indexing_stats().images_not_described == {"spec.pdf (in page-1): lost image 1"}
