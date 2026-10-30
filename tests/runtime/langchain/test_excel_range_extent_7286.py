# Copyright (c) 2026 EPAM Systems
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Merge, hyperlink and comment ranges never stretch an Excel sheet past its written cells (#7286)."""

import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from elitea_sdk.runtime.langchain.document_loaders import EliteAExcelLoader as excel_module
from elitea_sdk.runtime.langchain.document_loaders.EliteAExcelLoader import (
    EXCEL_MAX_REQUEST_ROWS,
    EliteAExcelLoader,
    ExcelReadLimitExceeded,
)

CONTACTS_EXAMPLE = str(Path(__file__).parent / "test_data" / "contacts_example.xlsx")
FIVE_ROWS = "name | value\nrow 1 | 1\nrow 2 | 2\nrow 3 | 3\nrow 4 | 4"


def five_row_workbook(tmp_path, name, sheet_xml=""):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["name", "value"])
    for index in range(1, 5):
        sheet.append([f"row {index}", index])
    plain = tmp_path / f"plain_{name}"
    workbook.save(plain)
    if not sheet_xml:
        return str(plain)
    annotated = tmp_path / name
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(annotated, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<pageMargins", sheet_xml.encode() + b"<pageMargins")
            target.writestr(item, data)
    return str(annotated)


WHOLE_COLUMN_RANGES = [
    pytest.param('<mergeCells count="1"><mergeCell ref="C1:C1048576"/></mergeCells>', id="merge"),
    pytest.param('<hyperlinks><hyperlink ref="D1:D1048576" location="Data!A1"/></hyperlinks>', id="hyperlink"),
]


def read_file(path, **kwargs):
    return EliteAExcelLoader(file_path=path, file_name=path, **kwargs).get_content()


def indexed(path, **kwargs):
    loader = EliteAExcelLoader(file_path=path, file_name=path, skip_size_check=True, **kwargs)
    return "\n".join(document.page_content for document in loader.load())


def read_with_formulas(path, monkeypatch):
    monkeypatch.setattr(EliteAExcelLoader, "_compute_formula_values", lambda self, file_path: {})
    return read_file(path, evaluate_formulas=True, max_tokens=-1)


@pytest.mark.parametrize("sheet_xml", WHOLE_COLUMN_RANGES)
def test_a_whole_column_range_adds_no_rows_when_indexing(tmp_path, sheet_xml):
    path = five_row_workbook(tmp_path, "whole_column.xlsx", sheet_xml)

    assert indexed(path, max_tokens=512) == FIVE_ROWS


@pytest.mark.parametrize("sheet_xml", WHOLE_COLUMN_RANGES)
def test_a_whole_column_range_adds_no_rows_to_read_file(tmp_path, sheet_xml):
    path = five_row_workbook(tmp_path, "whole_column.xlsx", sheet_xml)

    assert read_file(path, max_tokens=-1) == {"Data": [FIVE_ROWS]}


@pytest.mark.parametrize("sheet_xml", [
    pytest.param('<mergeCells count="1"><mergeCell ref="C1:C10000"/></mergeCells>', id="merge"),
    pytest.param('<hyperlinks><hyperlink ref="D1:D10000" location="Data!A1"/></hyperlinks>', id="hyperlink"),
])
def test_a_long_range_adds_no_rows_when_evaluating_formulas(tmp_path, monkeypatch, sheet_xml):
    path = five_row_workbook(tmp_path, "long_range.xlsx", sheet_xml)

    assert read_with_formulas(path, monkeypatch) == {"Data": [FIVE_ROWS]}


def test_a_whole_column_hyperlink_binds_no_cells_outside_the_sheet(tmp_path):
    path = five_row_workbook(tmp_path, "whole_column.xlsx", WHOLE_COLUMN_RANGES[1].values[0])
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        layout = excel_module._bind_streamed_sheet_layout(workbook.active)
    finally:
        workbook.close()

    assert layout.links == {}


def test_a_comment_far_below_the_data_adds_no_rows(tmp_path):
    workbook = Workbook()
    workbook.active.append(["a", "b"])
    workbook.active["Z5000"].comment = Comment("note", "author")
    saved = tmp_path / "saved.xlsx"
    workbook.save(saved)
    path = str(tmp_path / "far_comment.xlsx")
    with zipfile.ZipFile(saved) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<row r="5000"><c r="Z5000" t="n"></c></row>', b"")
            target.writestr(item, data)

    assert indexed(path, max_tokens=-1) == "a | b"
    assert read_file(path, max_tokens=-1) == {"Sheet": ["a | b"]}


def test_the_attached_contacts_workbook_keeps_its_rows_and_drops_the_link_column_tail():
    content = read_file(CONTACTS_EXAMPLE, max_tokens=-1)
    contacts = content["Contacts"][0].splitlines()

    assert list(content) == ["Contacts", "Directory"]
    assert len(contacts) == 136
    assert contacts[2].startswith(" | Anna Petrosyan | anna.p@example.com | EMEA | +374 10 000001 | [Open directory](None)")
    assert contacts[-1].startswith(" | Hana Sato | hana.s@example.com |  | +374 10 000134 | [Directory!A1](None)")
    assert content["Directory"] == ["Team | Lead\nSales | Anna Petrosyan\nSupport | Ben Carter"]


def test_the_attached_contacts_workbook_indexes_like_read_file():
    read = read_file(CONTACTS_EXAMPLE, max_tokens=512)

    assert indexed(CONTACTS_EXAMPLE, max_tokens=512) == "\n".join(chunk for chunks in read.values() for chunk in chunks)


def test_read_file_opens_a_plain_workbook_read_only(monkeypatch):
    opened = []
    real_load_workbook = excel_module.load_workbook
    monkeypatch.setattr(excel_module, "load_workbook",
                        lambda *args, **kwargs: opened.append(kwargs) or real_load_workbook(*args, **kwargs))

    read_file(CONTACTS_EXAMPLE, max_tokens=-1)

    assert opened and all(kwargs.get("read_only") for kwargs in opened)


def test_a_hyperlink_range_inside_the_data_still_fills_its_empty_cells(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    for row in [["h1", "h2", "h3"], ["a", "x", None], ["b", None, "y"], ["c", "w", "v"]]:
        sheet.append(row)
    sheet["B2"].hyperlink = "https://e.x/b"
    path = str(tmp_path / "in_data.xlsx")
    workbook.save(path)
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(tmp_path / "range.xlsx", "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="B2"', b'ref="B2:B4"')
            target.writestr(item, data)

    assert indexed(str(tmp_path / "range.xlsx"), max_tokens=-1).splitlines() == [
        "h1 | h2 | h3",
        "a | [x](https://e.x/b) | ",
        "b | [https://e.x/b](https://e.x/b) | y",
        "c | [w](https://e.x/b) | v",
    ]


def test_a_styled_empty_cell_below_the_data_still_counts_as_written(tmp_path):
    workbook = Workbook()
    workbook.active.append(["a", "b"])
    workbook.active["C4"].style = "Good"
    path = str(tmp_path / "styled.xlsx")
    workbook.save(path)

    assert indexed(path, max_tokens=-1) == "a | b | \n |  | \n |  | \n |  | "


def test_the_read_guard_still_rejects_an_oversized_read_through_load(tmp_path):
    workbook = Workbook()
    for index in range(EXCEL_MAX_REQUEST_ROWS + 10):
        workbook.active.append([f"row {index}"])
    path = str(tmp_path / "large.xlsx")
    workbook.save(path)

    with pytest.raises(ExcelReadLimitExceeded):
        EliteAExcelLoader(file_path=path, file_name=path).load()


def sheet_without_cells(tmp_path, name, sheet_xml=""):
    workbook = Workbook()
    workbook.active.title = "Empty"
    for row in range(1, 2_001):
        workbook.active.row_dimensions[row].height = 20
    saved = tmp_path / f"saved_{name}"
    workbook.save(saved)
    path = str(tmp_path / name)
    with zipfile.ZipFile(saved) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<pageMargins", sheet_xml.encode() + b"<pageMargins")
            target.writestr(item, data)
    return path


SHEETS_WITHOUT_CELLS = [
    pytest.param("", id="rows-with-a-height"),
    pytest.param('<mergeCells count="1"><mergeCell ref="A1:A10000"/></mergeCells>', id="long-merge"),
]


@pytest.mark.parametrize("sheet_xml", SHEETS_WITHOUT_CELLS)
def test_a_sheet_without_cells_reads_as_empty(tmp_path, monkeypatch, sheet_xml):
    path = sheet_without_cells(tmp_path, "no_cells.xlsx", sheet_xml)

    assert read_file(path, max_tokens=-1) == {"Empty": [""]}
    assert indexed(path, max_tokens=-1) == ""
    assert read_with_formulas(path, monkeypatch) == {"Empty": [""]}
