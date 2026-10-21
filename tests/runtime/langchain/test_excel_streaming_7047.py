# Copyright (c) 2026 EPAM Systems
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Indexing reads .xlsx in streaming read-only mode with the same output as the full read (#7047)."""

import datetime
import re
import tracemalloc
import zipfile

import pytest
import tiktoken
from langchain_core.documents import Document
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.hyperlink import Hyperlink

from elitea_sdk.runtime.langchain.document_loaders import EliteAExcelLoader as excel_module
from elitea_sdk.runtime.langchain.document_loaders.EliteAExcelLoader import (
    EXCEL_MAX_REQUEST_ROWS,
    EliteAExcelLoader,
)
from elitea_sdk.tools.utils import content_parser

ENCODING = tiktoken.get_encoding("cl100k_base")


def window_tokens(chunk):
    return sum(len(ENCODING.encode(row)) for row in chunk.splitlines())


def build_workbook(tmp_path, write_only=False):
    workbook = Workbook(write_only=write_only)
    if not write_only:
        workbook.remove(workbook.active)
    emails = workbook.create_sheet("Emails")
    emails.append(["Subject", "Sender", "Sent", "Body", "Score"])
    for index in range(40):
        emails.append([f"Subject {index}", None if index % 7 == 0 else f"sender{index}@example.com",
                       datetime.datetime(2026, 1, 1 + index % 28), "word " * (index % 13), index * 1.5])
    notes = workbook.create_sheet("Нотатки")
    notes.append(["Key", "Value"])
    notes.append(["alpha", 1])
    notes.append([None, None])
    notes.append(["omega", "last"])
    if not write_only:
        emails["A2"].hyperlink = "https://example.com/first"
        emails["B5"].hyperlink = "https://example.com/sender"
        emails["A9"].hyperlink = Hyperlink(ref="A9", location="'Нотатки'!A1")
    path = tmp_path / ("write_only.xlsx" if write_only else "streamed.xlsx")
    workbook.save(path)
    return str(path)


def loaded(path, **kwargs):
    loader = EliteAExcelLoader(file_path=path, file_name=path, **kwargs)
    return [(document.metadata["sheet_name"], document.page_content) for document in loader.load()]


LOADER_CONFIGS = [
    pytest.param({"max_tokens": -1}, id="one-chunk-per-sheet"),
    pytest.param({"max_tokens": 40}, id="token-windows"),
    pytest.param({"max_tokens": 40, "add_header_to_chunks": True}, id="header-per-window"),
    pytest.param({"max_tokens": 40, "add_header_to_chunks": True, "header_row_number": 2}, id="second-row-header"),
    pytest.param({"max_tokens": 40, "sheet_name": "Нотатки"}, id="named-sheet"),
    pytest.param({"max_tokens": 40, "sheet_name": "Missing"}, id="missing-sheet"),
]


@pytest.mark.parametrize("config", LOADER_CONFIGS)
def test_streamed_read_matches_the_full_read(tmp_path, monkeypatch, config):
    path = build_workbook(tmp_path)
    streamed = loaded(path, skip_size_check=True, **config)
    monkeypatch.setattr(EliteAExcelLoader, "_streams_workbook", lambda self: False)

    assert streamed == loaded(path, skip_size_check=True, **config)


def test_every_window_starts_with_the_header_exactly_once(tmp_path):
    path = build_workbook(tmp_path)

    windows = [chunk.splitlines() for sheet, chunk in
               loaded(path, skip_size_check=True, max_tokens=40, add_header_to_chunks=True) if sheet == "Emails"]

    assert len(windows) > 1
    assert all(lines[0].startswith("Subject | Sender") for lines in windows)
    assert all(sum(line.startswith("Subject | Sender") for line in lines) == 1 for lines in windows)


def test_streamed_read_renders_external_and_internal_hyperlinks(tmp_path):
    path = build_workbook(tmp_path)

    content = "\n".join(chunk for _, chunk in loaded(path, skip_size_check=True, max_tokens=-1))

    assert "[Subject 0](https://example.com/first)" in content
    assert "[sender3@example.com](https://example.com/sender)" in content
    assert "[Subject 7](None)" in content


def test_streamed_read_opens_the_workbook_read_only_and_closes_it(tmp_path, monkeypatch):
    path = build_workbook(tmp_path)
    opened = []
    real_load_workbook = excel_module.load_workbook

    def spy(*args, **kwargs):
        workbook = real_load_workbook(*args, **kwargs)
        opened.append((kwargs, workbook))
        return workbook

    monkeypatch.setattr(excel_module, "load_workbook", spy)

    loaded(path, skip_size_check=True, max_tokens=40)

    assert [kwargs.get("read_only") for kwargs, _ in opened] == [True]
    assert opened[0][1]._archive.fp is None


def test_streamed_read_keeps_every_row_when_the_dimension_tag_is_missing(tmp_path):
    path = build_workbook(tmp_path, write_only=True)

    rows = "\n".join(chunk for _, chunk in loaded(path, skip_size_check=True, max_tokens=-1)).splitlines()

    assert len(rows) == 41 + 4
    assert rows[-1] == "omega | last"


def test_formula_evaluation_keeps_the_full_read(tmp_path, monkeypatch):
    path = build_workbook(tmp_path)
    opened = []
    real_load_workbook = excel_module.load_workbook
    monkeypatch.setattr(excel_module, "load_workbook",
                        lambda *args, **kwargs: opened.append(kwargs) or real_load_workbook(*args, **kwargs))

    loaded(path, skip_size_check=True, evaluate_formulas=True, max_tokens=40)

    assert opened and not any(kwargs.get("read_only") for kwargs in opened)


def test_documents_are_produced_one_window_at_a_time(tmp_path):
    path = build_workbook(tmp_path)
    loader = EliteAExcelLoader(file_path=path, file_name=path, skip_size_check=True, max_tokens=40)

    documents = loader.iter_documents()
    first = next(documents)

    assert first.metadata["sheet_name"] == "Emails"
    assert window_tokens(first.page_content) <= 40
    documents.close()


def index_workbook(path, chunking_config):
    with open(path, "rb") as source:
        content = source.read()
    return list(content_parser.process_document_by_type(
        content=content, extension_source="emails.xlsx",
        document=Document(page_content="", metadata={"id": "emails"}), chunking_config=chunking_config))


def large_workbook(tmp_path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Large"
    sheet.append(["Subject", "Body"])
    for index in range(rows):
        sheet.append([f"Subject {index}", f"body text number {index}"])
    path = tmp_path / "large.xlsx"
    workbook.save(path)
    return str(path)


def test_indexing_keeps_one_chunk_per_sheet_the_form_submits_as_unbounded(tmp_path, monkeypatch):
    path = large_workbook(tmp_path, 2_000)
    form_config = {".xlsx": {"max_tokens": -1, "skip_size_check": False}}
    streamed = [chunk.page_content for chunk in index_workbook(path, form_config)]
    monkeypatch.setattr(EliteAExcelLoader, "_streams_workbook", lambda self: False)

    assert len(streamed) == 1
    assert streamed == [chunk.page_content for chunk in index_workbook(path, form_config)]


def test_indexing_keeps_a_token_window_the_user_chose(tmp_path):
    path = large_workbook(tmp_path, 2_000)

    chunks = index_workbook(path, {".xlsx": {"max_tokens": 64}})

    assert max(window_tokens(chunk.page_content) for chunk in chunks) <= 64


def test_indexing_reads_a_sheet_above_the_read_file_row_limit(tmp_path):
    path = large_workbook(tmp_path, EXCEL_MAX_REQUEST_ROWS + 10)

    chunks = index_workbook(path, {})

    assert not chunks[0].page_content.startswith("Error during content parsing")
    assert f"Subject {EXCEL_MAX_REQUEST_ROWS + 9} | body text number {EXCEL_MAX_REQUEST_ROWS + 9}" in chunks[-1].page_content


def test_loaders_that_stream_are_not_materialised():
    class StreamingLoader:
        def iter_documents(self):
            yield Document(page_content="streamed")

        def load(self):
            raise AssertionError("load must not be called for a streaming loader")

    assert [document.page_content for document in content_parser._iter_loaded_documents(StreamingLoader())] == ["streamed"]


def with_wrong_dimension_tags(path):
    rewritten = path.replace(".xlsx", "_wrong_dimension.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="A1"/>', data)
            target.writestr(item, data)
    return rewritten


def test_a_wrong_dimension_tag_does_not_cut_the_sheet(tmp_path, monkeypatch):
    path = with_wrong_dimension_tags(build_workbook(tmp_path))
    streamed = loaded(path, skip_size_check=True, max_tokens=-1)
    monkeypatch.setattr(EliteAExcelLoader, "_streams_workbook", lambda self: False)

    assert streamed == loaded(path, skip_size_check=True, max_tokens=-1)
    assert "[sender3@example.com](https://example.com/sender)" in streamed[0][1]


def test_an_empty_sheet_streams_like_the_full_read(tmp_path, monkeypatch):
    workbook = Workbook()
    workbook.active.title = "Empty"
    workbook.create_sheet("Filled").append(["a", "b"])
    path = str(tmp_path / "with_empty.xlsx")
    workbook.save(path)
    streamed = loaded(path, skip_size_check=True, max_tokens=40)
    monkeypatch.setattr(EliteAExcelLoader, "_streams_workbook", lambda self: False)

    assert streamed == loaded(path, skip_size_check=True, max_tokens=40)


def test_a_loader_failing_after_it_yielded_does_not_restart_on_a_fallback_extension(monkeypatch):
    class HalfwayLoader:
        def __init__(self, file_path=None, **kwargs):
            pass

        def iter_documents(self):
            yield Document(page_content="first window")
            raise MemoryError()

    monkeypatch.setitem(content_parser.loaders_map, ".xlsx", {
        "class": HalfwayLoader, "kwargs": {}, "allowed_to_override": {}})
    produced = []

    with pytest.raises(MemoryError):
        for document in content_parser.process_content_by_type(
                b"payload", "attachment.xlsx", fallback_extensions=[".txt"]):
            produced.append(document.page_content)

    assert produced == ["first window"]


def without_empty_cells(path, dropped_cells=(), moved_refs=(), added_merges=()):
    rewritten = path.replace(".xlsx", "_sparse.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(rb'<c r="[A-Z]+[0-9]+"(?: s="[0-9]+")?\s*/>', b"", data)
                for coordinate in dropped_cells:
                    data = re.sub(rb'<c r="%s"[^>]*>.*?</c>' % coordinate, b"", data)
                for old_ref, new_ref in moved_refs:
                    data = data.replace(b'ref="%s"' % old_ref, b'ref="%s"' % new_ref)
                data = re.sub(rb'<row [^>]*(?:/>|></row>)', b"", data)
                if added_merges:
                    merges = b"".join(b'<mergeCell ref="%s"/>' % ref for ref in added_merges)
                    data = data.replace(b"</sheetData>", b'</sheetData><mergeCells count="%d">%s</mergeCells>'
                                        % (len(added_merges), merges))
            target.writestr(item, data)
    return rewritten


def sparse_workbook(tmp_path, name, annotate):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Annotated"
    sheet["A1"], sheet["B1"] = "h1", "h2"
    sheet["A1"].hyperlink = "http://a.com"
    sheet["A2"] = "merged"
    sheet["A4"] = "m2"
    dropped_cells, moved_refs, *added_merges = annotate(sheet)
    path = str(tmp_path / f"{name}.xlsx")
    workbook.save(path)
    return without_empty_cells(path, dropped_cells, moved_refs, *added_merges)


def merge_beyond_the_cells(sheet):
    sheet.merge_cells("A2:D2")
    return [], []


def link_beyond_the_cells(sheet):
    sheet["E9"] = "link text Excel did not store"
    sheet["E9"].hyperlink = "http://empty.com"
    return [b"E9"], []


def comment_beyond_the_cells(sheet):
    sheet["G12"].comment = Comment("note", "author")
    return [b"G12"], []


def links_on_merged_cells(sheet):
    sheet.merge_cells("A2:D2")
    sheet["B6"] = "wide"
    sheet.merge_cells("B6:C7")
    linked = {"H20": "http://anchor.com", "J30": "http://range.com",
              "F3": Hyperlink(ref="F3", location="'Annotated'!A1")}
    for coordinate, link in linked.items():
        sheet[coordinate] = "link text Excel did not store"
        sheet[coordinate].hyperlink = link
    return [b"H20", b"J30", b"F3"], [(b"H20", b"C7"), (b"J30", b"A2:B4")]


def links_overlapping_an_empty_cell(sheet):
    sheet["E9"] = "link text Excel did not store"
    sheet["E9"].hyperlink = Hyperlink(ref="E9", tooltip="link with neither target nor location")
    sheet["H20"] = "link text Excel did not store"
    sheet["H20"].hyperlink = "http://second.com"
    return [b"E9", b"H20"], [(b"H20", b"E9")]


def merge_below_the_cells(sheet):
    sheet.merge_cells("A4:A7")
    return [], []


def value_in_a_hidden_merged_cell(sheet):
    sheet["B2"] = "hidden"
    sheet["C2"] = "z"
    return [], [], [b"A2:B2"]


SPARSE_LAYOUTS = [
    pytest.param(merge_beyond_the_cells, id="merge-beyond-cells"),
    pytest.param(link_beyond_the_cells, id="link-beyond-cells"),
    pytest.param(comment_beyond_the_cells, id="comment-beyond-cells"),
    pytest.param(links_on_merged_cells, id="links-on-merged-cells"),
    pytest.param(links_overlapping_an_empty_cell, id="links-overlapping-an-empty-cell"),
    pytest.param(merge_below_the_cells, id="merge-below-cells"),
    pytest.param(value_in_a_hidden_merged_cell, id="value-in-hidden-merged-cell"),
]


@pytest.mark.parametrize("annotate", SPARSE_LAYOUTS)
def test_the_sparse_fixture_omits_empty_cells_like_excel(tmp_path, annotate):
    with zipfile.ZipFile(sparse_workbook(tmp_path, annotate.__name__, annotate)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")

    written_cells = re.findall(rb'<c r="([A-Z]+[0-9]+)"', sheet_xml)
    assert sorted(written_cells) == sorted([b"A1", b"B1", b"A2", b"A4"] + {
        links_on_merged_cells: [b"B6"], value_in_a_hidden_merged_cell: [b"B2", b"C2"]}.get(annotate, []))
    assert re.findall(rb'<row r="([0-9]+)"', sheet_xml) == sorted(
        {re.match(rb"[A-Z]+([0-9]+)", cell).group(1) for cell in written_cells}, key=int)


@pytest.mark.parametrize("config", [
    pytest.param({"max_tokens": -1}, id="one-chunk-per-sheet"),
    pytest.param({"max_tokens": 8}, id="token-windows"),
])
@pytest.mark.parametrize("annotate", SPARSE_LAYOUTS)
def test_merges_links_and_comments_outside_the_written_cells_stream_like_the_full_read(
        tmp_path, monkeypatch, annotate, config):
    path = sparse_workbook(tmp_path, annotate.__name__, annotate)
    streamed = loaded(path, skip_size_check=True, **config)
    monkeypatch.setattr(EliteAExcelLoader, "_streams_workbook", lambda self: False)

    assert streamed == loaded(path, skip_size_check=True, **config)


def test_a_whole_column_merge_costs_memory_per_range_not_per_cell(tmp_path):
    workbook = Workbook()
    workbook.active.append(["a", "b"])
    path = str(tmp_path / "column_merge.xlsx")
    workbook.save(path)
    path = without_empty_cells(path, added_merges=[b"C1:C1048576"])
    read_only = load_workbook(path, read_only=True, data_only=True)
    try:
        tracemalloc.start()
        layout = excel_module._bind_streamed_sheet_layout(read_only.active)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
        read_only.close()

    assert read_only.active.max_row == 1_048_576
    assert peak < 1_000_000
    assert excel_module._render_streamed_row(("a", "b", "x"), 2, layout) == "a | b | "


def report_with_a_merge_and_a_link_per_row(tmp_path, rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, rows + 1):
        sheet.append([f"item {row}", f"merged {row}", None, None, f"link {row}"])
        sheet.merge_cells(f"B{row}:C{row}")
        sheet[f"E{row}"].hyperlink = f"http://example.com/{row}"
    sheet["G1"] = "link on a hidden merged cell"
    sheet["G1"].hyperlink = "http://example.com/hidden"
    path = str(tmp_path / "report.xlsx")
    workbook.save(path)
    return without_empty_cells(path, dropped_cells=[b"G1"], moved_refs=[(b"G1", b"C1")])


def test_placing_links_checks_only_the_merges_on_their_row(tmp_path, monkeypatch):
    rows = 3_000
    path = report_with_a_merge_and_a_link_per_row(tmp_path, rows)
    checks = []
    hides = excel_module._MergedRange.hides

    def counted_hides(self, position):
        checks.append(position)
        return hides(self, position)

    read_only = load_workbook(path, read_only=True, data_only=True)
    try:
        monkeypatch.setattr(excel_module._MergedRange, "hides", counted_hides)
        layout = excel_module._bind_streamed_sheet_layout(read_only.active)
    finally:
        read_only.close()

    assert len(checks) <= 2 * (rows + 1)
    assert layout.links[(1, 2)].target == "http://example.com/hidden"


def test_a_report_with_a_merge_and_a_link_per_row_streams_like_the_full_read(tmp_path, monkeypatch):
    path = report_with_a_merge_and_a_link_per_row(tmp_path, 50)
    streamed = loaded(path, skip_size_check=True, max_tokens=-1)
    monkeypatch.setattr(EliteAExcelLoader, "_streams_workbook", lambda self: False)

    assert streamed == loaded(path, skip_size_check=True, max_tokens=-1)
